"""
Pipeline de dados do SalesPulse.

Lê o Excel consolidado (FatoVendas_Consolidado + Dim_Consultores) e
regenera o arquivo `salespulse/src/data.json` consumido pelo dashboard.

Uso (dentro do .venv ativado):
    python atualizar_dados.py
"""
import datetime
import json
import os
import sys

import openpyxl

# terminal Windows (cp1252) não imprime emojis/acentos por padrão
try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

PASTA = os.path.dirname(os.path.abspath(__file__))
EXCEL = os.path.join(
    PASTA, "PowerBI_Vendas_Maio_Junho_Consolidado.xlsx"
)
SAIDA = os.path.join(PASTA, "salespulse", "src", "data.json")

METAS = {"Nathy": 5200.0, "Gisele": 5000.0, "Kalyne": 5000.0}


def linhas(ws):
    it = ws.iter_rows(values_only=True)
    next(it)  # pula cabeçalho
    return [dict(zip(next(ws.iter_rows(values_only=True)), r)) for r in it]


def normalizar_data(d):
    if isinstance(d, datetime.datetime):
        return d.strftime("%Y-%m-%d")
    return str(d).split("T")[0][:10]


def main():
    wb = openpyxl.load_workbook(EXCEL, read_only=True)

    consultores, _ = [], []
    for r in list(wb["Dim_Consultores"].iter_rows(values_only=True))[1:]:
        consultores.append(
            {"id": r[0], "nome": r[1], "equipe": r[2]}
        )

    vendas = []
    for r in list(wb["FatoVendas_Consolidado"].iter_rows(values_only=True))[1:]:
        vendas.append(
            {
                "data": normalizar_data(r[0]),
                "consultorId": r[1],
                "valor": round(float(r[2]), 2),
            }
        )

    # preserva metas editadas manualmente no data.json, se existirem
    metas = METAS
    if os.path.exists(SAIDA):
        try:
            atual = json.load(open(SAIDA, encoding="utf-8"))
            metas = atual.get("metas", METAS)
        except Exception:
            pass

    out = {"consultores": consultores, "vendas": vendas, "metas": metas}
    json.dump(out, open(SAIDA, "w", encoding="utf-8"), ensure_ascii=False, indent=1)

    print(f"✅ {SAIDA}")
    print(f"   {len(consultores)} consultores | {len(vendas)} transações")
    print(f"   metas: {metas}")


if __name__ == "__main__":
    main()
